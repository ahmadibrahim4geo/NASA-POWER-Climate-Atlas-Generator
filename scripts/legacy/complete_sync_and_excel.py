# -*- coding: utf-8 -*-
"""
complete_sync_and_excel.py
--------------------------
Completes Step 5 (Excel updates with Seasonal_Range) and syncs all rasters, GDB layers,
and tables to D: Drive.
"""
from __future__ import print_function
import os
import sys
import shutil
import io
import openpyxl
import arcpy

DESKTOP_DIR = r"C:\Users\ahmad\Desktop\Egypt Climate Data 1996-2025"
DDRIVE_DIR = r"D:\MY DATA & FILE\Spatial Data\Egypt Climate Data 1996-2025"

DESKTOP_GDB = os.path.join(DESKTOP_DIR, "Climate_Database_From_1996_To_2025.gdb")
DDRIVE_GDB = os.path.join(DDRIVE_DIR, "Climate_Database_From_1996_To_2025.gdb")

SEASONAL_RANGE_DEFS = [
    ("Temperature", "T_Seasonal_Range", u"المدى الفصلي لدرجة الحرارة",
     ["T_Winter_Mean", "T_Spring_Mean", "T_Summer_Mean", "T_Autumn_Mean"],
     os.path.join("01_Temperature", "T_Seasonal_Range")),
    ("Sea_Level_Pressure", "PSL_Seasonal_Range", u"المدى الفصلي لضغط مستوى البحر",
     ["PSL_Winter_Mean", "PSL_Spring_Mean", "PSL_Summer_Mean", "PSL_Autumn_Mean"],
     os.path.join("03_Sea_Level_Pressure", "PSL_Seasonal_Range")),
    ("Surface_Pressure", "PS_Seasonal_Range", u"المدى الفصلي للضغط السطحي",
     ["PS_Winter_Mean", "PS_Spring_Mean", "PS_Summer_Mean", "PS_Autumn_Mean"],
     os.path.join("04_Surface_Pressure", "PS_Seasonal_Range")),
    ("Wind", "W_Spd_Seasonal_Range", u"المدى الفصلي لسرعة الرياح",
     ["W_Spd_Winter_Mean", "W_Spd_Spring_Mean", "W_Spd_Summer_Mean", "W_Spd_Autumn_Mean"],
     os.path.join("05_Wind", "Speed", "W_Spd_Seasonal_Range")),
    ("Relative_Humidity", "RH_Seasonal_Range", u"المدى الفصلي للرطوبة النسبية",
     ["RH_Winter_Mean", "RH_Spring_Mean", "RH_Summer_Mean", "RH_Autumn_Mean"],
     os.path.join("06_Relative_Humidity", "RH_Seasonal_Range")),
    ("Dew_Point", "Td_Seasonal_Range", u"المدى الفصلي لنقطة الندى",
     ["Td_Winter_Mean", "Td_Spring_Mean", "Td_Summer_Mean", "Td_Autumn_Mean"],
     os.path.join("07_Dew_Point", "Td_Seasonal_Range")),
    ("Solar_Radiation", "Sol_Seasonal_Range", u"المدى الفصلي للإشعاع الشمسي",
     ["Sol_Winter_Mean", "Sol_Spring_Mean", "Sol_Summer_Mean", "Sol_Autumn_Mean"],
     os.path.join("08_Solar_Radiation", "Sol_Seasonal_Range")),
    ("UV_Index", "UV_Seasonal_Range", u"المدى الفصلي للأشعة فوق البنفسجية",
     ["UV_Winter_Mean", "UV_Spring_Mean", "UV_Summer_Mean", "UV_Autumn_Mean"],
     os.path.join("09_UV_Index", "UV_Seasonal_Range")),
    ("Cloud_Cover", "Cld_Seasonal_Range", u"المدى الفصلي للغطاء السحابي",
     ["Cld_Winter_Mean", "Cld_Spring_Mean", "Cld_Summer_Mean", "Cld_Autumn_Mean"],
     os.path.join("10_Cloud_Cover", "Cld_Seasonal_Range")),
]

ALIAS_FIX_MAP = {
    "ET_January_Total": u"مجموع البخر-نتح لشهر يناير (ملم)",
    "ET_February_Total": u"مجموع البخر-نتح لشهر فبراير (ملم)",
    "ET_March_Total": u"مجموع البخر-نتح لشهر مارس (ملم)",
    "ET_April_Total": u"مجموع البخر-نتح لشهر أبريل (ملم)",
    "ET_May_Total": u"مجموع البخر-نتح لشهر مايو (ملم)",
    "ET_June_Total": u"مجموع البخر-نتح لشهر يونيو (ملم)",
    "ET_July_Total": u"مجموع البخر-نتح لشهر يوليو (ملم)",
    "ET_August_Total": u"مجموع البخر-نتح لشهر أغسطس (ملم)",
    "ET_September_Total": u"مجموع البخر-نتح لشهر سبتمبر (ملم)",
    "ET_October_Total": u"مجموع البخر-نتح لشهر أكتوبر (ملم)",
    "ET_November_Total": u"مجموع البخر-نتح لشهر نوفمبر (ملم)",
    "ET_December_Total": u"مجموع البخر-نتح لشهر ديسمبر (ملم)",
    "Wind_Speed": u"سرعة الرياح (متر/ث)",
    "Wind_Dir": u"اتجاه الرياح (درجة)",
    "Arrow_Angle": u"زاوية دوران السهم الكارتوجرافي",
    "Arrow_Size": u"حجم السهم النسبي",
    "Period": u"الفترة المناخية",
    "Id": u"معرف خط الضغط المتساوي",
    "Contour": u"قيمة خط الضغط الجوي المتساوي (hPa)",
    "NAME": u"اسم النطاق الجغرافي",
}

def update_excel_workbooks(base_dir):
    print(">>> Updating Excel workbooks in %s..." % base_dir)
    tables_dir = os.path.join(base_dir, "00_Tables_And_Reports")
    gdb = os.path.join(base_dir, "Climate_Database_From_1996_To_2025.gdb")

    excel_map = {
        "Temperature": ("Temperature.xls", "T_Seasonal_Range", "T_Annual_Range"),
        "Sea_Level_Pressure": ("Sea_Level_Pressure.xls", "PSL_Seasonal_Range", "PSL_Annual_Range"),
        "Surface_Pressure": ("Surface_Pressure.xls", "PS_Seasonal_Range", "PS_Annual_Range"),
        "Wind": ("Wind.xls", "W_Spd_Seasonal_Range", "W_Spd_Annual_Range"),
        "Relative_Humidity": ("Relative_Humidity.xls", "RH_Seasonal_Range", "RH_Annual_Range"),
        "Dew_Point": ("Dew_Point.xls", "Td_Seasonal_Range", "Td_Annual_Range"),
        "Solar_Radiation": ("Solar_Radiation.xls", "Sol_Seasonal_Range", "Sol_Annual_Range"),
        "UV_Index": ("UV_Index.xls", "UV_Seasonal_Range", "UV_Annual_Range"),
        "Cloud_Cover": ("Cloud_Cover.xls", "Cld_Seasonal_Range", "Cld_Annual_Range"),
    }

    for fc, (xls_name, range_fld, after_fld) in excel_map.items():
        xls_path = os.path.join(tables_dir, xls_name)
        if not os.path.isfile(xls_path):
            print("  Warning: %s not found" % xls_path)
            continue
        
        pts_fc = os.path.join(gdb, fc)
        pt_values = {}
        with arcpy.da.SearchCursor(pts_fc, ["OBJECTID", range_fld]) as cur:
            for oid, val in cur:
                pt_values[int(oid)] = val

        with open(xls_path, "rb") as f:
            buf = io.BytesIO(f.read())
        wb = openpyxl.load_workbook(buf)
        if "Data" not in wb.sheetnames:
            continue
        
        ws = wb["Data"]
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        if range_fld in header:
            print("  %s already contains %s." % (xls_name, range_fld))
            continue
        
        insert_col = ws.max_column + 1
        if after_fld in header:
            insert_col = header.index(after_fld) + 2
        
        ws.insert_cols(insert_col)
        ws.cell(1, insert_col).value = range_fld

        oid_col = header.index("OBJECTID") + 1 if "OBJECTID" in header else 1
        for r in range(2, ws.max_row + 1):
            oid_val = ws.cell(r, oid_col).value
            try:
                oid = int(oid_val)
            except Exception:
                oid = r - 1
            ws.cell(r, insert_col).value = pt_values.get(oid, 0.0)

        wb.save(xls_path)
        print("  Updated %s with column %s at col %d." % (xls_name, range_fld, insert_col))

def sync_to_ddrive():
    print(">>> Syncing updates to D: Drive...")
    
    # 1. Fix aliases on D: Drive GDB
    arcpy.env.workspace = DDRIVE_GDB
    for fc in arcpy.ListFeatureClasses():
        for f in arcpy.ListFields(fc):
            if f.name in ALIAS_FIX_MAP:
                try:
                    arcpy.management.AlterField(fc, f.name, new_field_alias=ALIAS_FIX_MAP[f.name])
                except Exception:
                    pass

    # 2. Add and compute Seasonal Range on D: Drive GDB
    for fc, range_fld, alias_ar, season_flds, rel_base in SEASONAL_RANGE_DEFS:
        existing = [f.name for f in arcpy.ListFields(fc)]
        if range_fld not in existing:
            arcpy.management.AddField(fc, range_fld, "DOUBLE", field_alias=alias_ar)
        else:
            try:
                arcpy.management.AlterField(fc, range_fld, new_field_alias=alias_ar)
            except Exception:
                pass
        
        fld_list = season_flds + [range_fld]
        with arcpy.da.UpdateCursor(fc, fld_list) as cur:
            for row in cur:
                vals = [row[i] for i in range(4) if row[i] is not None]
                if len(vals) == 4:
                    row[4] = round(float(max(vals) - min(vals)), 2)
                    cur.updateRow(row)
        print("  D: Drive GDB: updated %s in %s." % (range_fld, fc))

    # 3. Copy generated rasters (.tif, .tfw, .aux.xml, .ovr, .xml) from Desktop to D: Drive
    for fc, range_fld, alias_ar, season_flds, rel_base in SEASONAL_RANGE_DEFS:
        src_tif = os.path.join(DESKTOP_DIR, rel_base + ".tif")
        dst_tif = os.path.join(DDRIVE_DIR, rel_base + ".tif")
        dst_dir = os.path.dirname(dst_tif)
        if not os.path.exists(dst_dir):
            os.makedirs(dst_dir)
        
        src_parent = os.path.dirname(src_tif)
        base_name = os.path.basename(src_tif)
        for fname in os.listdir(src_parent):
            if fname.startswith(os.path.splitext(base_name)[0]):
                shutil.copy2(os.path.join(src_parent, fname), os.path.join(dst_dir, fname))
        print("  Copied raster %s to D: Drive." % base_name)

    # 4. Copy Vector feature classes from Desktop GDB to D: Drive GDB
    vector_fcs = [
        "Wind_Vector_Month",
        "Isobars_PSL_Month_Mean",
        "Isobars_PS_Month_Mean",
        "Isobars_PSL_Annual_Mean",
        "Isobars_PS_Annual_Mean"
    ]
    for vfc in vector_fcs:
        src_vfc = os.path.join(DESKTOP_GDB, vfc)
        dst_vfc = os.path.join(DDRIVE_GDB, vfc)
        if arcpy.Exists(src_vfc):
            if arcpy.Exists(dst_vfc):
                try: arcpy.management.Delete(dst_vfc)
                except Exception: pass
            arcpy.management.CopyFeatures(src_vfc, dst_vfc)
            print("  Copied feature class %s to D: Drive GDB." % vfc)

    # 5. Copy shapefiles for Wind_Vector_Month
    src_shp_dir = os.path.join(DESKTOP_DIR, "05_Wind", "Direction", "Vector_Points")
    dst_shp_dir = os.path.join(DDRIVE_DIR, "05_Wind", "Direction", "Vector_Points")
    if os.path.exists(src_shp_dir):
        if not os.path.exists(dst_shp_dir):
            os.makedirs(dst_shp_dir)
        for f in os.listdir(src_shp_dir):
            if f.startswith("Wind_Vector_Month"):
                shutil.copy2(os.path.join(src_shp_dir, f), os.path.join(dst_shp_dir, f))
        print("  Copied Wind_Vector_Month shapefile to D: Drive.")

    # 6. Copy updated Excel workbooks from Desktop to D: Drive
    src_xls_dir = os.path.join(DESKTOP_DIR, "00_Tables_And_Reports")
    dst_xls_dir = os.path.join(DDRIVE_DIR, "00_Tables_And_Reports")
    for f in os.listdir(src_xls_dir):
        if f.endswith(".xls") or f.endswith(".xlsx"):
            shutil.copy2(os.path.join(src_xls_dir, f), os.path.join(dst_xls_dir, f))
    print("  Copied all updated Excel workbooks to D: Drive.")

def main():
    print("==================================================")
    print("EXECUTING EXCEL UPDATES & D: DRIVE SYNCHRONIZATION")
    print("==================================================")
    update_excel_workbooks(DESKTOP_DIR)
    sync_to_ddrive()
    print("==================================================")
    print("ALL SYNCHRONIZATION COMPLETED!")
    print("==================================================")

if __name__ == "__main__":
    main()
