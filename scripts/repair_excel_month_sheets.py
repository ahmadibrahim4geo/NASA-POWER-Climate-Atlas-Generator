# -*- coding: utf-8 -*-
"""
Repair element workbooks: monthly Jan-Dec columns live ONLY in the Month
sheet, never in Data. Also fixes the kPa/mbar unit bug left by an earlier
manual backfill (Month sheets held raw kPa while Data holds mbar).

- XLSX element files: strip monthly cols from Data; overwrite Month values
  from the computed climatology table (keyed by OBJECTID).
- Precipitation.xls (BIFF): strip Data tail; (re)build Month sheet.
- Climate_Atlas_Master_Workbook.xls: strip monthly tails from all sheets.
"""
import io
import os
import shutil
import sys
import tempfile

from importlib.machinery import SourceFileLoader

BASE = r"C:\Users\ahmad\Desktop\Egypt Climate Data 1996-2025"
TABLES = os.path.join(BASE, "00_Tables_And_Reports")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
BASE_COLS = ["OBJECTID", "Source_ID", "Point_Lat", "Point_Lon"]


def load_tool():
    return SourceFileLoader(
        "pytmod", os.path.join(ROOT, "POWER_Climate_Atlas_Generator_10_8.pyt")).load_module()


def is_monthly_header(h):
    return isinstance(h, str) and any(m in h for m in MONTHS)


def monthly_fields_of(pyt, module):
    return [f for f in pyt.MODULE_FIELDS.get(module, []) if pyt.is_monthly_field(f)]


def compute_table(pyt):
    """Recompute the 12-month climatology from the raw JSON archive."""
    import json
    raw = json.load(open(os.path.join(TABLES, "Egypt_Monthly_Raw_1996_2025.json"),
                         "r", encoding="utf-8-sig"))
    coords = {}
    import openpyxl
    tmp = os.path.join(tempfile.gettempdir(), "opencode_rc.xlsx")
    shutil.copy(os.path.join(TABLES, "Temperature.xls"), tmp)
    wb = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
    sh = wb["Data"]
    rows = list(sh.iter_rows(values_only=True))
    hdr = rows[0]
    for r in rows[1:]:
        if r[0] is None:
            continue
        coords[str(int(r[hdr.index("OBJECTID")]))] = float(r[hdr.index("Point_Lat")])
    years = list(range(1996, 2026))
    mods = ["Temperature", "Precipitation", "Sea Level Pressure",
            "Surface Pressure", "Wind", "Relative Humidity", "Dew Point",
            "Solar Radiation", "UV Index", "Cloud Cover", "Evapotranspiration"]
    table = {}
    for pid, series in raw.items():
        monthly = {}
        for param, s in series.items():
            monthly[param] = {}
            for k, v in s.items():
                try:
                    m = int(k[4:])
                except Exception:
                    continue
                if 1 <= m <= 12:
                    monthly[param][(int(k[:4]), m)] = v
        table[pid] = pyt.compute_point_fields(
            monthly, years, mods, "Monthly", lat=coords.get(pid, 25.0))
    return table


def repair_xlsx_element(path, module, flds, table):
    import openpyxl
    tmp = os.path.join(tempfile.gettempdir(), "opencode_fx.xlsx")
    shutil.copy(path, tmp)
    wb = openpyxl.load_workbook(tmp)
    # --- Data: drop monthly columns ---
    ds = wb["Data"]
    hdr = [c.value for c in next(ds.iter_rows(min_row=1, max_row=1))]
    cut = next((i for i, h in enumerate(hdr) if is_monthly_header(h)), None)
    if cut is not None:
        ds.delete_cols(cut + 1, ds.max_column - cut)
    # --- Month: ensure headers, overwrite values ---
    if "Month" in wb.sheetnames:
        ms = wb["Month"]
    else:
        ms = wb.create_sheet("Month")
    want = BASE_COLS + flds
    ms.delete_cols(1, max(ms.max_column, 1))
    for j, h in enumerate(want):
        ms.cell(row=1, column=j + 1, value=h)
    dhdr = [c.value for c in next(ds.iter_rows(min_row=1, max_row=1))]
    i_oid = dhdr.index("OBJECTID")
    i_sid = dhdr.index("Source_ID")
    i_lat = dhdr.index("Point_Lat") if "Point_Lat" in dhdr else None
    i_lon = dhdr.index("Point_Lon") if "Point_Lon" in dhdr else None
    r_out = 2
    for row in ds.iter_rows(min_row=2, values_only=True):
        if row[i_oid] is None:
            continue
        pid = str(int(row[i_oid]))
        vals = table.get(pid, {})
        ms.cell(row=r_out, column=1, value=row[i_oid])
        ms.cell(row=r_out, column=2, value=row[i_sid])
        ms.cell(row=r_out, column=3, value=row[i_lat] if i_lat is not None else None)
        ms.cell(row=r_out, column=4, value=row[i_lon] if i_lon is not None else None)
        for j, f in enumerate(flds):
            ms.cell(row=r_out, column=5 + j, value=vals.get(f))
        r_out += 1
    wb.save(tmp)
    shutil.copy(tmp, path)
    print("repaired XLSX: %s (Data %d cols, Month %d rows)" % (
        os.path.basename(path), ds.max_column, r_out - 2))


def rewrite_biff(pyt, path, sheets):
    """sheets: [(name, header, rows)]. Styled, like the tool writer."""
    pyt.write_master_excel_workbook(path, [(n, h, r, False) for n, h, r in sheets])


def repair_precipitation(pyt, table):
    import xlrd
    path = os.path.join(TABLES, "Precipitation.xls")
    wb = xlrd.open_workbook(path)
    sh = wb.sheet_by_index(0)
    hdr = [sh.cell_value(0, c) for c in range(sh.ncols)]
    cut = next((i for i, h in enumerate(hdr) if is_monthly_header(h)), len(hdr))
    data_hdr = hdr[:cut]
    data_rows = [[sh.cell_value(r, c) for c in range(cut)] for r in range(1, sh.nrows)]
    flds = monthly_fields_of(pyt, "Precipitation")
    m_hdr = BASE_COLS + flds
    i_oid = data_hdr.index("OBJECTID")
    m_rows = []
    for r in data_rows:
        vals = table.get(str(int(r[i_oid])), {})
        m_rows.append([r[data_hdr.index("OBJECTID")], r[data_hdr.index("Source_ID")],
                       r[data_hdr.index("Point_Lat")], r[data_hdr.index("Point_Lon")]] +
                      [vals.get(f) for f in flds])
    rewrite_biff(pyt, path, [("Data", data_hdr, data_rows), ("Month", m_hdr, m_rows)])
    print("repaired BIFF: Precipitation.xls (Data %d, Month %d rows)" % (len(data_hdr), len(m_rows)))


def repair_master(pyt):
    import xlrd
    path = os.path.join(TABLES, "Climate_Atlas_Master_Workbook.xls")
    wb = xlrd.open_workbook(path)
    sheets = []
    for s in wb.sheet_names():
        sh = wb.sheet_by_name(s)
        hdr = [sh.cell_value(0, c) for c in range(sh.ncols)]
        cut = next((i for i, h in enumerate(hdr) if is_monthly_header(h)), len(hdr))
        if cut < len(hdr):
            print("  master/%s: stripped %d monthly cols" % (s, len(hdr) - cut))
        hdr = hdr[:cut]
        rows = [[sh.cell_value(r, c) for c in range(cut)] for r in range(1, sh.nrows)]
        sheets.append((s, hdr, rows))
    rewrite_biff(pyt, path, sheets)
    print("repaired BIFF: Climate_Atlas_Master_Workbook.xls")


XLSX_ELEMENTS = {
    "Temperature.xls": "Temperature",
    "Sea_Level_Pressure.xls": "Sea Level Pressure",
    "Surface_Pressure.xls": "Surface Pressure",
    "Wind.xls": "Wind",
    "Relative_Humidity.xls": "Relative Humidity",
    "Dew_Point.xls": "Dew Point",
    "Solar_Radiation.xls": "Solar Radiation",
    "UV_Index.xls": "UV Index",
    "Cloud_Cover.xls": "Cloud Cover",
    "Evapotranspiration.xls": "Evapotranspiration",
}


def main():
    pyt = load_tool()
    table = compute_table(pyt)
    for fname, module in XLSX_ELEMENTS.items():
        repair_xlsx_element(os.path.join(TABLES, fname), module,
                            monthly_fields_of(pyt, module), table)
    repair_precipitation(pyt, table)
    repair_master(pyt)
    print("DONE.")


if __name__ == "__main__":
    main()
